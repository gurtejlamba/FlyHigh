"""Generate SIFT presentation by filling the HYPNEXIS26 official template."""
import copy
import shutil
from pathlib import Path
from lxml import etree
from pptx import Presentation

TEMPLATE = Path("/Users/gurtej/Downloads/HYPNEXIS26_Official_PPT_Template.pptx")
OUTPUT   = Path("docs/SIFT_presentation.pptx")

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ── helpers ──────────────────────────────────────────────────────────────────

def _txBody(shape):
    return shape.text_frame._txBody


def set_text(shape, text: str):
    """Replace all text in a shape with *text*, preserving run formatting."""
    tb = _txBody(shape)
    paras = tb.findall(f"{{{NS_A}}}p")
    if not paras:
        return
    # Keep only the first paragraph; delete the rest
    for p in paras[1:]:
        tb.remove(p)
    p0 = paras[0]
    runs = p0.findall(f"{{{NS_A}}}r")
    if not runs:
        return
    # Keep first run, delete extras
    for r in runs[1:]:
        p0.remove(r)
    runs[0].find(f"{{{NS_A}}}t").text = text


def set_bullets(shape, lines: list[str]):
    """Replace bullet paragraphs with *lines*, preserving paragraph formatting."""
    tb = _txBody(shape)
    paras = tb.findall(f"{{{NS_A}}}p")
    if not paras:
        return
    template_para = paras[0]
    # Remove all existing bullet paragraphs
    for p in paras:
        tb.remove(p)
    for line in lines:
        new_p = copy.deepcopy(template_para)
        runs = new_p.findall(f"{{{NS_A}}}r")
        if runs:
            for r in runs[1:]:
                new_p.remove(r)
            runs[0].find(f"{{{NS_A}}}t").text = line
        tb.append(new_p)


def find(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


# ── content ──────────────────────────────────────────────────────────────────

TEAM_TAG = "AI-02  |  H-AI02"

SLIDES = {
    # ── slide 1: title ──────────────────────────────────────────────────────
    1: {
        "Text 6":  "AI-02",
        "Text 8":  "H-AI02",
        "Text 10": "PS-AI-02",
        "Text 12": "Data-Analyst Agent on Messy Databases",
        "Text 14": "Artificial Intelligence / NLP",
        "Text 16": "Software",
    },

    # ── slide 2: problem & solution ─────────────────────────────────────────
    2: {
        "Text 0":   "SIFT — SQL Intelligence For Tables",
        "Team tag": TEAM_TAG,
        "Problem body": [
            "Real databases are polluted: NULLs, 'N/A' strings, mixed date formats",
            "Analysts waste hours writing SQL for ad-hoc questions on live data",
            "LLM-based tools hallucinate table names and fail on messy schemas",
        ],
        "Solution body": [
            "4-stage pipeline: Explore → Generate → Verify → Abstain",
            "Self-repairs SQL on execution failure (up to 3 retries with error feedback)",
            "Abstains cleanly rather than returning a confidently wrong answer",
        ],
        "Image placeholder text": "SIFT Pipeline: Explore → Generate → Verify → Abstain",
    },

    # ── slide 3: technical approach ─────────────────────────────────────────
    3: {
        "Text 0":   "TECHNICAL APPROACH",
        "Team tag": TEAM_TAG,
        "Arch body": [
            "Explore: PRAGMA table_info + DISTINCT samples → compact SchemaInfo",
            "Generate: schema-aware prompt → Qwen2.5-Coder-7B → raw SQL",
            "Verify: thread-timeout execution + NULL-ratio guard on results",
            "Abstain: hard-error regex → returns structured 'cannot answer'",
        ],
        "Tech body": [
            "Python 3.13 · sqlite3 · threading (timeout isolation)",
            "Qwen2.5-Coder-7B-Instruct via Ollama (fully local, zero cost)",
            "BIRD benchmark · python-pptx · reportlab · pytest (37 tests)",
        ],
        "Unique body": [
            "Iterative self-repair: error text fed back as prompt context on retry",
            "Schema column quoting, CAST rules, ORDER BY anti-pattern rules baked into prompt",
            "Runs entirely offline — no API key, no cloud, no data egress",
        ],
    },

    # ── slide 4: UX & design ────────────────────────────────────────────────
    4: {
        "Text 0":   "DEMO & WORKFLOW",
        "Team tag": TEAM_TAG,
        "UI body": [
            "CLI: python -m src.agent.loop <db> '<question>'",
            "Python API: loop.run(db_path, question) → RunResult",
            "JSON output with SQL, rows, latency, attempt count",
        ],
        "Text 9":  "User asks a natural-language question about their database",
        "Text 12": "SIFT explores the schema, generates SQL, and executes it",
        "Text 15": "Verified answer returned; or a clean abstain with reason",
        "Goal body": [
            "Any analyst can query any SQLite database in plain English, "
            "with guaranteed-safe execution and honest abstention.",
        ],
        "Wireframe placeholder": "[ CLI demo: python -m src.agent.loop california_schools.sqlite ]",
    },

    # ── slide 5: feasibility ────────────────────────────────────────────────
    5: {
        "Text 0":   "FEASIBILITY AND VIABILITY",
        "Team tag": TEAM_TAG,
        "Technical Feasibility body": [
            "Runs on 8 GB VRAM or CPU-only (Qwen 7B quantised)",
            "Read-only SQLite: zero risk of data corruption",
            "Modular stages — swap the LLM or verifier independently",
        ],
        "Economic Feasibility body": [
            "Zero inference cost: fully local, no API fees",
            "BIRD benchmark eval suite included for continuous measurement",
            "Scales to larger models (13B / 34B) with same codebase",
        ],
        "Operational Viability body": [
            "Thread-timeout prevents hung queries; abstain prevents wrong answers",
            "EX = 0.60 on 20-question BIRD sample (+25 pts over baseline 0.35)",
            "37 unit tests cover explore, generate, verify, abstain, and eval",
        ],
    },

    # ── slide 6: impact ─────────────────────────────────────────────────────
    6: {
        "Text 0":   "IMPACT AND BENEFITS",
        "Team tag": TEAM_TAG,
        "Social Impact body": [
            "Democratises data access for non-technical staff",
            "Reduces dependency on specialised SQL engineers",
            "Enables faster decisions from live operational databases",
        ],
        "Environmental Impact body": [
            "Fully local: no cloud compute, no data-centre emissions",
            "7B model fits a single consumer GPU — minimal energy footprint",
            "No network round-trips; lower latency and carbon per query",
        ],
        "Technical Impact body": [
            "Advances Text-to-SQL robustness on messy real-world schemas",
            "Prompt-engineering techniques for CAST, quoting, ORDER BY anti-patterns",
            "Open evaluation harness on BIRD benchmark for reproducibility",
        ],
        "Educational Impact body": [
            "Demonstrates schema introspection, LLM prompt design, and evaluation",
            "Reusable pipeline for teaching AI-assisted data analysis",
            "Full test suite and eval runner serve as learning reference",
        ],
    },

    # ── slide 7: business model ─────────────────────────────────────────────
    7: {
        "Text 0":   "BUSINESS MODEL & SUSTAINABILITY",
        "Team tag": TEAM_TAG,
        "Users body": [
            "Data analysts in SMEs with local SQLite / DuckDB databases",
            "Government agencies requiring air-gapped data processing",
            "Research labs running structured experiment logs",
        ],
        "Rev body": [
            "Open-source core (MIT); professional support tier",
            "Enterprise licence for on-premise deployment & SLA",
            "Consulting: schema design + prompt fine-tuning packages",
        ],
        "Sus body": [
            "Community contributions via GitHub; plugin architecture",
            "Model-agnostic design: upgrades free as open-source models improve",
            "BIRD leaderboard submissions keep the project visible & benchmarked",
        ],
    },

    # ── slide 8: references ─────────────────────────────────────────────────
    8: {
        "Text 0":   "RESEARCH AND REFERENCES",
        "Team tag": TEAM_TAG,
        "Ref body": [
            "Li et al. 'Can LLM Already Serve as a Database Interface?' NeurIPS 2023 (BIRD benchmark)",
            "Qwen Team. 'Qwen2.5-Coder Technical Report.' Alibaba Cloud, 2024",
            "Ollama — local LLM inference engine: https://ollama.ai",
            "python-pptx, reportlab, pytest — open-source Python libraries",
            "SQLite 3.45 — embedded relational database: https://sqlite.org",
        ],
    },

    # ── slide 9: thank you ──────────────────────────────────────────────────
    9: {
        "Text 1": "AI-02  |  H-AI02  |  lambagurtejsingh@gmail.com",
    },
}


# ── main ─────────────────────────────────────────────────────────────────────

def build():
    prs = Presentation(str(TEMPLATE))

    for slide_idx, replacements in SLIDES.items():
        slide = prs.slides[slide_idx - 1]
        for shape_name, content in replacements.items():
            shape = find(slide, shape_name)
            if shape is None:
                print(f"  WARNING: shape '{shape_name}' not found on slide {slide_idx}")
                continue
            if isinstance(content, list):
                set_bullets(shape, content)
            else:
                set_text(shape, content)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    import os
    os.chdir(Path(__file__).parent.parent)
    build()
